"""
השוואת תמלול מקצועי - AlefBot מול OpenAI, במקביל
==================================================
דף /compare: מעלים קובץ אודיו/וידאו, או מדביקים קישור להקלטה קיימת (הקישור
"להורדת ההקלטה" מהמייל של תמלול פון, או כל קישור ישיר לקובץ). שני המנועים
רצים במקביל על אותו קובץ בדיוק, ובסוף:
  - נשלח מייל עם שני התמלולים זה לצד זה + זמן עיבוד, אורך, אחוז דמיון,
    ושני קובצי Word מצורפים.
  - התוצאה נשמרת ומוצגת בדף /compare/<id> (זה לצד זה), ויש רשימת השוואות קודמות.

AlefBot - בדיוק כמו "תמלול מקצועי" במערכת הראשית: model_tier='standard',
translate_to_hebrew לפי בחירת שפת הפלט.
OpenAI - מודל התמלול (ברירת מחדל gpt-transcribe, אפשר לשנות במשתנה
LAB_OPENAI_TRANSCRIBE_MODEL). הקובץ מומר ל-WAV 16kHz מונו ומחולק לחלקים של
10 דקות (מגבלת גודל/אורך של OpenAI), כל חלק מתומלל, והטקסט מחובר לפי הסדר.

עלות (אופציונלי): אם מוגדרים LAB_PRICE_ALEFBOT_PER_MIN / LAB_PRICE_OPENAI_PER_MIN
(בשקלים לדקה) - מוצגת הערכת עלות לכל מנוע לפי אורך ההקלטה.

נפרד לגמרי מהמערכת הראשית: לא נוגע ב-DB שלה, בלקוחות או בחיובים.
"""
import os
import io
import re
import time
import uuid
import base64
import sqlite3
import difflib
import logging
import tempfile
import threading
import subprocess
from datetime import datetime
from html import escape

import requests
from flask import request, render_template_string, redirect

log = logging.getLogger(__name__)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
COMPARE_DIR = os.path.join(BASE_DIR, 'compare_files')
os.makedirs(COMPARE_DIR, exist_ok=True)
COMPARE_DB = os.path.join(BASE_DIR, 'compare.db')

ALEFBOT_BASE = 'https://alef-bot.top/api/v1'
OPENAI_MODEL = os.environ.get('LAB_OPENAI_TRANSCRIBE_MODEL', 'gpt-transcribe')
CHUNK_SECONDS = 600
MAX_FILE_SIZE = 300 * 1024 * 1024
OPENAI_PROMPT = 'תמלול שיחה או שיעור בעברית, לעתים עם ארמית ומינוח תורני, ראשי תיבות וציטוטים.'

LANG_NAMES = {'he': 'עברית', 'yi': 'אידיש', 'en': 'אנגלית'}
ENGINE_NAMES = {'alefbot': 'AlefBot (התמלול המקצועי היום)', 'openai': f'OpenAI ({OPENAI_MODEL})'}


# ---------------------------------------------------------------- DB

def _db():
    conn = sqlite3.connect(COMPARE_DB)
    conn.row_factory = sqlite3.Row
    conn.execute("""CREATE TABLE IF NOT EXISTS comparisons (
        id TEXT PRIMARY KEY, created_at TEXT, source TEXT, email TEXT,
        language TEXT, output_language TEXT, duration_sec REAL,
        alefbot_status TEXT, alefbot_text TEXT, alefbot_error TEXT, alefbot_secs REAL,
        openai_status TEXT, openai_text TEXT, openai_error TEXT, openai_secs REAL,
        emailed INTEGER DEFAULT 0
    )""")
    return conn


def _update(cid, **fields):
    with _lock:
        conn = _db()
        cols = ', '.join(f'{k} = ?' for k in fields)
        conn.execute(f'UPDATE comparisons SET {cols} WHERE id = ?', (*fields.values(), cid))
        conn.commit()
        conn.close()


def _get(cid):
    conn = _db()
    row = conn.execute('SELECT * FROM comparisons WHERE id = ?', (cid,)).fetchone()
    conn.close()
    return dict(row) if row else None


def _list(limit=50):
    conn = _db()
    rows = conn.execute('SELECT id, created_at, source, alefbot_status, openai_status, duration_sec '
                        'FROM comparisons ORDER BY created_at DESC LIMIT ?', (limit,)).fetchall()
    conn.close()
    return [dict(r) for r in rows]


_lock = threading.Lock()


# ---------------------------------------------------------------- אודיו

def _ffmpeg():
    try:
        import imageio_ffmpeg
        return imageio_ffmpeg.get_ffmpeg_exe()
    except Exception:
        return 'ffmpeg'


def _to_wav16k(src, dst):
    """ממיר כל קובץ אודיו/וידאו ל-WAV 16kHz מונו 16 ביט. מחזיר אורך בשניות."""
    subprocess.run([_ffmpeg(), '-y', '-v', 'error', '-i', src, '-vn', '-ac', '1', '-ar', '16000',
                    '-sample_fmt', 's16', dst], check=True, timeout=1800)
    import wave
    with wave.open(dst, 'rb') as w:
        return w.getnframes() / float(w.getframerate())


def _split_wav(wav_path, out_dir):
    """מחלק WAV לחלקים של CHUNK_SECONDS (בלי לטעון את כל הקובץ לזיכרון)."""
    import wave
    parts = []
    with wave.open(wav_path, 'rb') as w:
        params = w.getparams()
        per = params.framerate * CHUNK_SECONDS
        n = 0
        while True:
            data = w.readframes(per)
            if not data:
                break
            p = os.path.join(out_dir, f'part_{n:03d}.wav')
            with wave.open(p, 'wb') as o:
                o.setparams(params)
                o.writeframes(data)
            parts.append(p)
            n += 1
    return parts


# ---------------------------------------------------------------- מנועים

def _run_alefbot(wav_path, output_language):
    """אותן הגדרות כמו התמלול המקצועי במערכת הראשית. מחזיר טקסט או זורק חריגה."""
    key = os.environ.get('ALEFBOT_API_KEY')
    if not key:
        raise RuntimeError('לא הוגדר ALEFBOT_API_KEY')
    h = {'Authorization': f'Bearer {key}'}
    size = os.path.getsize(wav_path)
    up = requests.post(f'{ALEFBOT_BASE}/uploads', headers={**h, 'Content-Type': 'application/json'},
                       json={'filename': f'compare_{uuid.uuid4().hex[:8]}.wav', 'content_type': 'audio/wav',
                             'size_bytes': size}, timeout=30)
    up.raise_for_status()
    upload_id = up.json().get('upload_id')
    with open(wav_path, 'rb') as fh:
        requests.put(f'{ALEFBOT_BASE}/uploads/{upload_id}/binary', headers=h, data=fh, timeout=900).raise_for_status()
    tr = requests.post(f'{ALEFBOT_BASE}/transcriptions', headers={**h, 'Content-Type': 'application/json'},
                       json={'upload_id': upload_id, 'output_format': 'plain_text', 'model_tier': 'standard',
                             'translate_to_hebrew': output_language == 'he'}, timeout=30)
    tr.raise_for_status()
    job_id = tr.json().get('job_id') or tr.json().get('id')
    if not job_id:
        raise RuntimeError('AlefBot לא החזיר מספר עבודה')
    for _ in range(240):   # עד שעתיים
        time.sleep(30)
        try:
            st = requests.get(f'{ALEFBOT_BASE}/transcriptions/{job_id}', headers=h, timeout=15).json().get('status', '')
        except Exception as e:
            log.warning(f'compare alefbot poll: {e}')
            continue
        if st == 'completed':
            art = requests.get(f'{ALEFBOT_BASE}/transcriptions/{job_id}/artifact?format=txt', headers=h, timeout=60)
            art.raise_for_status()
            return art.text.strip()
        if st in ('failed', 'cancelled'):
            raise RuntimeError(f'AlefBot: העבודה {st}')
    raise RuntimeError('AlefBot: לא הסתיים אחרי שעתיים')


def _run_openai(wav_path, language):
    from openai import OpenAI
    client = OpenAI(api_key=os.environ.get('OPENAI_API_KEY'))
    tmp = tempfile.mkdtemp(prefix='cmp_parts_', dir=COMPARE_DIR)
    try:
        texts = []
        for p in _split_wav(wav_path, tmp):
            last = None
            for attempt in range(3):
                try:
                    with open(p, 'rb') as fh:
                        kw = {'model': OPENAI_MODEL, 'file': (os.path.basename(p), fh, 'audio/wav'),
                              'response_format': 'text', 'prompt': OPENAI_PROMPT}
                        if language in ('he', 'en', 'yi'):
                            kw['language'] = language
                        res = client.audio.transcriptions.create(**kw)
                    texts.append((res if isinstance(res, str) else getattr(res, 'text', '') or '').strip())
                    last = None
                    break
                except Exception as e:
                    last = e
                    log.warning(f'compare openai part {p} attempt {attempt + 1}: {e}')
                    time.sleep(5)
            if last:
                raise RuntimeError(f'OpenAI נכשל בחלק {os.path.basename(p)}: {last}')
        return '\n\n'.join(t for t in texts if t)
    finally:
        import shutil
        shutil.rmtree(tmp, ignore_errors=True)


# ---------------------------------------------------------------- תהליך

def _process(cid, src_path, link):
    work = tempfile.mkdtemp(prefix='cmp_', dir=COMPARE_DIR)
    try:
        if link:
            src_path = os.path.join(work, 'source')
            with requests.get(link, stream=True, timeout=300) as r:
                r.raise_for_status()
                with open(src_path, 'wb') as fh:
                    for chunk in r.iter_content(1024 * 1024):
                        fh.write(chunk)
        wav = os.path.join(work, 'audio16k.wav')
        dur = _to_wav16k(src_path, wav)
        _update(cid, duration_sec=dur)
        row = _get(cid)

        def engine(name, fn):
            t0 = time.time()
            try:
                text = fn()
                _update(cid, **{f'{name}_status': 'done', f'{name}_text': text, f'{name}_secs': time.time() - t0})
            except Exception as e:
                log.error(f'compare {cid} {name}: {e}', exc_info=True)
                _update(cid, **{f'{name}_status': 'error', f'{name}_error': f'{type(e).__name__}: {e}',
                                f'{name}_secs': time.time() - t0})

        threads = [threading.Thread(target=engine, args=('alefbot', lambda: _run_alefbot(wav, row['output_language']))),
                   threading.Thread(target=engine, args=('openai', lambda: _run_openai(wav, row['language'])))]
        for t in threads:
            t.start()
        for t in threads:
            t.join()
    except Exception as e:
        log.error(f'compare {cid} prepare: {e}', exc_info=True)
        msg = f'הכנת הקובץ נכשלה: {type(e).__name__}: {e}'
        _update(cid, alefbot_status='error', alefbot_error=msg, openai_status='error', openai_error=msg)
    finally:
        import shutil
        shutil.rmtree(work, ignore_errors=True)
        if src_path and not link and os.path.exists(src_path):
            os.remove(src_path)
    try:
        _send_email(_get(cid))
        _update(cid, emailed=1)
    except Exception as e:
        log.error(f'compare {cid} email: {e}', exc_info=True)


# ---------------------------------------------------------------- מדדים ומייל

def similarity(a, b):
    """אחוז דמיון בין שני התמלולים לפי מילים (בלי פיסוק)."""
    wa = re.findall(r'\w+', a or '')
    wb = re.findall(r'\w+', b or '')
    if not wa and not wb:
        return None
    return round(difflib.SequenceMatcher(None, wa, wb, autojunk=False).ratio() * 100, 1)


def _price(name, dur):
    v = os.environ.get(f'LAB_PRICE_{name.upper()}_PER_MIN')
    try:
        return round(float(v) * (dur or 0) / 60.0, 2) if v else None
    except ValueError:
        return None


def summary(row):
    out = {}
    for n in ('alefbot', 'openai'):
        text = row.get(f'{n}_text') or ''
        out[n] = {'status': row.get(f'{n}_status'), 'text': text, 'error': row.get(f'{n}_error'),
                  'secs': row.get(f'{n}_secs'), 'words': len(re.findall(r'\w+', text)),
                  'chars': len(text), 'price': _price(n, row.get('duration_sec'))}
    out['similarity'] = similarity(out['alefbot']['text'], out['openai']['text'])
    return out


def _fmt_secs(s):
    if s is None:
        return '-'
    s = int(round(s))
    return f'{s // 60}:{s % 60:02d}'


def _docx(title, text):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    doc = Document()
    for p_text in [title] + (text or '').split('\n'):
        p = doc.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        bidi = OxmlElement('w:bidi')
        pPr.append(bidi)
        run = p.add_run(p_text)
        rPr = run._r.get_or_add_rPr()
        rtl = OxmlElement('w:rtl')
        rPr.append(rtl)
        if p_text is title:
            run.bold = True
            b = OxmlElement('w:bCs')
            rPr.append(b)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _send_email(row):
    import sendgrid
    from sendgrid.helpers.mail import Mail, Email, Attachment, FileContent, FileName, FileType, Disposition
    if not row or not row.get('email'):
        return
    s = summary(row)
    def col(n):
        e = s[n]
        body = (escape(e['text']) if e['status'] == 'done'
                else f'<span style="color:#b91c1c">שגיאה: {escape(e["error"] or "")}</span>')
        price = f' · עלות משוערת ₪{e["price"]:.2f}' if e['price'] is not None else ''
        return (f'<td style="vertical-align:top;width:50%;padding:12px;border:1px solid #cfe0fb;background:#ffffff">'
                f'<div style="font-weight:bold;color:#0a58ca;font-size:16px">{escape(ENGINE_NAMES[n])}</div>'
                f'<div style="font-size:13px;color:#43546b;margin:4px 0 10px">זמן עיבוד {_fmt_secs(e["secs"])} · '
                f'{e["words"]} מילים{price}</div>'
                f'<div style="white-space:pre-wrap;line-height:1.8;font-size:15px;text-align:justify">{body}</div></td>')
    sim = f'{s["similarity"]}%' if s['similarity'] is not None else '-'
    link = os.environ.get('APP_BASE_URL', '').rstrip('/')
    view = f'<p><a href="{link}/compare/{row["id"]}">לצפייה בהשוואה בדף</a> (צריך קוד גישה)</p>' if link else ''
    html = (f'<div dir="rtl" style="font-family:Arial,sans-serif;max-width:1000px;margin:0 auto;background:#eef4fb;padding:16px">'
            f'<h2 style="color:#0a58ca;margin:0 0 6px">השוואת תמלול: AlefBot מול OpenAI</h2>'
            f'<div style="color:#43546b;font-size:14px">{escape(row["source"] or "")} · אורך {_fmt_secs(row["duration_sec"])} · '
            f'שפה {LANG_NAMES.get(row["language"], row["language"])} · דמיון בין התמלולים {sim}</div>{view}'
            f'<table role="presentation" width="100%" cellpadding="0" cellspacing="0" style="border-collapse:collapse;margin-top:12px">'
            f'<tr>{col("alefbot")}{col("openai")}</tr></table></div>')
    msg = Mail(from_email=Email(os.environ.get('SENDGRID_FROM_EMAIL', ''), 'מעבדת בדיקות'),
               to_emails=row['email'], subject=f'🧪 השוואת תמלול - {row["source"]}', html_content=html)
    for n in ('alefbot', 'openai'):
        if s[n]['status'] == 'done' and s[n]['text']:
            try:
                data = _docx(f'{ENGINE_NAMES[n]} - {row["source"]}', s[n]['text'])
                msg.add_attachment(Attachment(FileContent(base64.b64encode(data).decode()), FileName(f'{n}.docx'),
                                              FileType('application/vnd.openxmlformats-officedocument.wordprocessingml.document'),
                                              Disposition('attachment')))
            except Exception as e:
                log.warning(f'compare docx {n}: {e}')
    sendgrid.SendGridAPIClient(api_key=os.environ.get('SENDGRID_API_KEY')).send(msg)


# ---------------------------------------------------------------- דפים

PAGE_HTML = """<!doctype html><html dir="rtl" lang="he"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>השוואת תמלול</title>
{% if refresh %}<meta http-equiv="refresh" content="20">{% endif %}
<style>
body{font-family:Arial,sans-serif;max-width:1100px;margin:30px auto;padding:0 16px;color:#0b1f3a;background:#eef4fb}
h2{color:#0a58ca;margin:0 0 6px}
.card{background:#fff;border:1px solid #cfe0fb;border-radius:12px;padding:18px;margin-top:16px}
label{display:block;margin-top:14px;font-weight:bold;font-size:14px}
input,select{width:100%;padding:10px;margin-top:4px;box-sizing:border-box;border:1px solid #cfe0fb;border-radius:8px;font-size:15px}
button{margin-top:18px;padding:12px 28px;background:#0a58ca;color:#fff;border:none;border-radius:24px;font-size:16px;cursor:pointer}
.note{color:#43546b;font-size:13px}
.err{color:#991b1b;background:#fef2f2;border:1px solid #fecaca;padding:10px;border-radius:8px;margin-top:12px}
.ok{color:#065f46;background:#f0fdf4;border:1px solid #bbf7d0;padding:10px;border-radius:8px;margin-top:12px}
.cols{display:grid;grid-template-columns:1fr 1fr;gap:14px;margin-top:14px}
@media (max-width:760px){.cols{grid-template-columns:1fr}}
.col{background:#fff;border:1px solid #cfe0fb;border-radius:12px;padding:14px;min-width:0}
.col h3{margin:0;color:#0a58ca;font-size:17px}
.meta{font-size:13px;color:#43546b;margin:4px 0 10px}
.text{white-space:pre-wrap;line-height:1.85;font-size:15px;text-align:justify}
table{width:100%;border-collapse:collapse;font-size:14px}
td,th{border-bottom:1px solid #e3ebf6;padding:8px;text-align:right}
a{color:#0a58ca}
.or{text-align:center;color:#43546b;margin-top:10px;font-size:13px}
</style></head><body>
<h2>השוואת תמלול: AlefBot מול OpenAI</h2>
<p class="note">שני המנועים רצים במקביל על אותו קובץ. התוצאה נשלחת למייל וגם מוצגת כאן. <a href="/">חזרה למעבדה</a></p>
{% if message %}<div class="{{ 'ok' if ok else 'err' }}">{{ message }}</div>{% endif %}

{% if row %}
  <div class="card">
    <b>{{ row.source }}</b>
    <div class="meta">נוצר {{ row.created_at }} · אורך {{ fmt(row.duration_sec) }} · שפה {{ langs.get(row.language, row.language) }}
      {% if s.similarity is not none %} · דמיון בין התמלולים {{ s.similarity }}%{% endif %}
      {% if refresh %} · <b>עדיין רץ - הדף מתרענן לבד</b>{% endif %}</div>
  </div>
  <div class="cols">
  {% for n in ['alefbot', 'openai'] %}{% set e = s[n] %}
    <div class="col"><h3>{{ names[n] }}</h3>
      <div class="meta">{% if e.status == 'done' %}זמן עיבוד {{ fmt(e.secs) }} · {{ e.words }} מילים · {{ e.chars }} תווים{% if e.price is not none %} · עלות משוערת ₪{{ '%.2f'|format(e.price) }}{% endif %}
        {% elif e.status == 'error' %}<span style="color:#b91c1c">שגיאה: {{ e.error }}</span>{% else %}רץ...{% endif %}</div>
      <div class="text">{{ e.text }}</div>
    </div>
  {% endfor %}
  </div>
{% else %}
  <form class="card" method="post" action="/compare/run" enctype="multipart/form-data">
    <input type="hidden" name="access_code" value="{{ access_code }}">
    <label>קובץ אודיו / וידאו</label>
    <input type="file" name="file" accept="audio/*,video/*">
    <div class="or">או</div>
    <label>קישור להקלטה קיימת</label>
    <input type="url" name="link" placeholder="למשל הקישור 'להורדת ההקלטה' מהמייל של תמלול פון">
    <p class="note">אפשר להדביק את הקישור מכל מייל תמלול, או כל קישור ישיר לקובץ אודיו.</p>
    <label>שפת ההקלטה</label>
    <select name="language"><option value="he">עברית</option><option value="yi">אידיש</option><option value="en">אנגלית</option></select>
    <label>שפת הפלט (ל-AlefBot)</label>
    <select name="output_language"><option value="he">עברית (תרגום לעברית)</option><option value="original">שפת ההקלטה</option></select>
    <label>שליחת ההשוואה למייל</label>
    <input type="email" name="email" value="{{ default_email }}" required>
    <button type="submit">הרץ השוואה</button>
  </form>
  {% if history %}
  <div class="card"><b>השוואות קודמות</b>
    <table><tr><th>תאריך</th><th>קובץ</th><th>אורך</th><th>AlefBot</th><th>OpenAI</th></tr>
    {% for h in history %}<tr><td>{{ h.created_at }}</td>
      <td><a href="/compare/{{ h.id }}?access_code={{ access_code }}">{{ h.source }}</a></td>
      <td>{{ fmt(h.duration_sec) }}</td><td>{{ st.get(h.alefbot_status, h.alefbot_status) }}</td><td>{{ st.get(h.openai_status, h.openai_status) }}</td></tr>{% endfor %}
    </table></div>
  {% endif %}
{% endif %}
</body></html>"""

STATUS_NAMES = {'pending': 'רץ', 'done': 'הסתיים', 'error': 'שגיאה'}

GATE_HTML = """<!doctype html><html dir="rtl" lang="he"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1"><title>השוואת תמלול</title>
<style>body{font-family:Arial,sans-serif;max-width:400px;margin:80px auto;padding:0 16px;color:#0b1f3a}
input{width:100%;padding:10px;margin-top:8px;box-sizing:border-box;border:1px solid #cfe0fb;border-radius:8px;font-size:15px}
button{margin-top:16px;padding:10px 24px;background:#0a58ca;color:#fff;border:none;border-radius:24px;cursor:pointer}
.err{color:#991b1b;background:#fef2f2;padding:10px;border-radius:8px;margin-top:12px}</style></head><body>
<h2>השוואת תמלול</h2>
<form method="get" action="{{ action }}"><label>קוד גישה</label>
<input type="password" name="access_code" required autofocus><button type="submit">כניסה</button></form>
{% if wrong %}<div class="err">קוד גישה שגוי</div>{% endif %}
</body></html>"""


def init_transcompare(app, *, require_access, default_email):

    def render(**kw):
        base = dict(message=None, ok=True, row=None, s=None, refresh=False, history=None, access_code='',
                    default_email=default_email, fmt=_fmt_secs, names=ENGINE_NAMES, langs=LANG_NAMES, st=STATUS_NAMES)
        base.update(kw)
        return render_template_string(PAGE_HTML, **base)

    def gate():
        code = request.values.get('access_code', '')
        if not code:
            return render_template_string(GATE_HTML, wrong=False, action=request.path), 200
        if not require_access(code):
            return render_template_string(GATE_HTML, wrong=True, action=request.path), 403
        return None

    @app.route('/compare', methods=['GET'])
    def compare_page():
        g = gate()
        if g:
            return g
        return render(access_code=request.args['access_code'], history=_list())

    @app.route('/compare/<cid>', methods=['GET'])
    def compare_view(cid):
        g = gate()
        if g:
            return g
        row = _get(cid)
        if not row:
            return render(access_code=request.args['access_code'], message='ההשוואה לא נמצאה', ok=False,
                          history=_list()), 404
        running = 'pending' in (row['alefbot_status'], row['openai_status'])
        return render(access_code=request.args['access_code'], row=row, s=summary(row), refresh=running)

    @app.route('/compare/run', methods=['POST'])
    def compare_run():
        code = request.form.get('access_code', '')
        if not require_access(code):
            return 'קוד גישה שגוי', 403
        f = request.files.get('file')
        link = (request.form.get('link') or '').strip()
        email = (request.form.get('email') or '').strip()
        language = request.form.get('language', 'he')
        output_language = request.form.get('output_language', 'he')
        has_file = bool(f and f.filename)
        if has_file == bool(link):
            return render(access_code=code, history=_list(), ok=False,
                          message='יש לבחור קובץ או להדביק קישור (אחד מהשניים)'), 400
        if link and not re.match(r'^https?://', link):
            return render(access_code=code, history=_list(), ok=False, message='הקישור צריך להתחיל ב-http'), 400
        cid = uuid.uuid4().hex[:12]
        src_path = None
        if has_file:
            f.seek(0, os.SEEK_END)
            size = f.tell()
            f.seek(0)
            if size > MAX_FILE_SIZE:
                return render(access_code=code, history=_list(), ok=False,
                              message=f'הקובץ גדול מדי ({size // 1024 // 1024}MB, מקסימום 300MB)'), 400
            ext = os.path.splitext(f.filename)[1][:10]
            src_path = os.path.join(COMPARE_DIR, f'{cid}{ext}')
            f.save(src_path)
            source = f.filename
        else:
            source = 'קישור: ' + (link.split('?')[0][-60:])
        with _lock:
            conn = _db()
            conn.execute('INSERT INTO comparisons (id, created_at, source, email, language, output_language, '
                         'alefbot_status, openai_status) VALUES (?, ?, ?, ?, ?, ?, ?, ?)',
                         (cid, datetime.now().strftime('%d/%m/%Y %H:%M'), source, email, language,
                          output_language, 'pending', 'pending'))
            conn.commit()
            conn.close()
        threading.Thread(target=_process, args=(cid, src_path, link or None), daemon=True).start()
        return redirect(f'/compare/{cid}?access_code={code}')
